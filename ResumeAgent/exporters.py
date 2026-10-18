"""ATS-friendly single-column DOCX export and optional offline PDF conversion."""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _append_text(paragraph, raw: str) -> None:
    # Basic Markdown inline bold. No icons, tables or text boxes for ATS readability.
    for i, part in enumerate(re.split(r"\*\*(.*?)\*\*", raw)):
        if not part:
            continue
        run = paragraph.add_run(part)
        if i % 2:
            run.bold = True


def to_docx(markdown_resume: str) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.62)
    sec.bottom_margin = Inches(.62)
    sec.left_margin = Inches(.77)
    sec.right_margin = Inches(.77)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(30, 40, 50)
    normal.paragraph_format.space_after = Pt(2.5)
    for label, size, before, after in [("Title", 17, 0, 5), ("Heading 2", 10.8, 9, 3)]:
        st = styles[label]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(21, 57, 82)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    lines = [s.rstrip() for s in markdown_resume.splitlines()]
    first_line = True
    for line in lines:
        clean = line.strip()
        if not clean:
            continue
        if clean.startswith("\\#"):  # escaped literal '#', not a heading
            p = doc.add_paragraph(style="Normal")
            _append_text(p, clean[1:])
            continue
        if clean.startswith("# "):
            p = doc.add_paragraph(style="Title")
            _append_text(p, clean[2:])
            first_line = False
        elif clean.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            _append_text(p, clean[3:])
        elif re.match(r"^[-*]\s+", clean):
            level = min(3, (len(line) - len(line.lstrip(" "))) // 2)
            p = doc.add_paragraph(style="Normal")
            p.style = styles["List Bullet"]
            p.paragraph_format.left_indent = Inches(.19 + .22 * level)
            p.paragraph_format.first_line_indent = Inches(-.12)
            p.paragraph_format.space_after = Pt(1.6)
            _append_text(p, re.sub(r"^[-*]\s+", "", clean))
        else:
            p = doc.add_paragraph(style="Normal")
            if first_line:
                p.style = styles["Title"]
                first_line = False
            _append_text(p, clean)
    # Metadata stays generic, not company-specific.
    doc.core_properties.title = "Tailored Resume"
    doc.core_properties.subject = "ATS-readable job-specific resume"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def to_pdf(docx_bytes: bytes) -> bytes:
    """Offline conversion: MS Word COM on Windows, otherwise local LibreOffice.

    A converter MUST be installed. We never use online conversion websites.
    """
    with tempfile.TemporaryDirectory(prefix="resume_agent_") as dirname:
        folder = Path(dirname)
        docx_path, pdf_path = folder / "resume.docx", folder / "resume.pdf"
        docx_path.write_bytes(docx_bytes)
        if sys.platform == "win32":
            try:
                import win32com.client
                app = win32com.client.DispatchEx("Word.Application")
                app.Visible = False
                app.DisplayAlerts = 0
                try:
                    document = app.Documents.Open(str(docx_path.resolve()), ReadOnly=True)
                    try:
                        document.ExportAsFixedFormat(str(pdf_path.resolve()), 17)
                    finally:
                        document.Close(False)
                finally:
                    app.Quit()
                if pdf_path.exists() and pdf_path.stat().st_size > 100:
                    return pdf_path.read_bytes()
            except Exception as exc:
                word_error = str(exc)
            else:
                word_error = "Word did not create a PDF"
        else:
            word_error = "Word COM is only available on Windows"
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice and sys.platform == "win32":
            for path in [r"C:\Program Files\LibreOffice\program\soffice.exe",
                         r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
                if Path(path).exists():
                    soffice = path
                    break
        if soffice:
            cmd = [soffice, "-env:UserInstallation=file:///" + str(folder / "profile").replace("\\", "/"),
                   "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(docx_path)]
            proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                  text=True, timeout=80, check=False)
            if proc.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 100:
                return pdf_path.read_bytes()
            raise RuntimeError("Local LibreOffice conversion failed: " + proc.stderr[-400:])
        raise RuntimeError("DOCX created. For PDF export, install Microsoft Word+pywin32 or LibreOffice. " + word_error[:250])


def cover_letter_docx(markdown_letter: str) -> bytes:
    """Professional, editable, single-column letter with genuine paragraphs."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.78)
    section.bottom_margin = Inches(.78)
    section.left_margin = Inches(.88)
    section.right_margin = Inches(.88)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 43, 52)
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.line_spacing = 1.12
    heading = doc.styles["Title"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(22, 56, 83)
    heading.paragraph_format.space_after = Pt(5)

    # Preserve paragraph breaks, not every displayed wrapped line.
    blocks = re.split(r"\n\s*\n", markdown_letter.strip())
    if len(blocks) == 1:
        blocks = [x for x in markdown_letter.splitlines() if x.strip()]
    for idx, block in enumerate(blocks):
        value = " ".join(x.strip() for x in block.splitlines() if x.strip())
        if not value:
            continue
        if value.startswith("# "):
            p = doc.add_paragraph(style="Title")
            _append_text(p, value[2:])
        else:
            p = doc.add_paragraph(style="Normal")
            # avoid literal Markdown ## in the letter; model is instructed otherwise
            _append_text(p, re.sub(r"^#{1,3}\s+", "", value))
            if value.lower().startswith("dear "):
                p.paragraph_format.space_after = Pt(13)
            if value.lower() in ("sincerely,", "best regards,", "kind regards,", "yours sincerely,"):
                p.paragraph_format.space_before = Pt(10)
    doc.core_properties.title = "Tailored Cover Letter"
    doc.core_properties.subject = "Role-specific application cover letter"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def to_text(markdown: str) -> str:
    """Plain text for pasting into online application forms."""
    out = []
    for line in markdown.splitlines():
        line = re.sub(r"\*\*(.*?)\*\*", r"\1", line.rstrip())
        if line.startswith("\\#"):
            line = line[1:]
        elif line.startswith("## "):
            line = "\n" + line[3:].upper()
        elif line.startswith("# "):
            line = line[2:]
        elif re.match(r"^\s*[-*]\s+", line):
            line = "• " + re.sub(r"^\s*[-*]\s+", "", line)
        out.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip() + "\n"


README_TEXT = ("Locally generated documents. Read the audit before applying.\n"
               "Scores are heuristic, not official ATS measurements.\n"
               "Validate identity, dates, achievements, employer names, and attachments.\n"
               "PDFs are included only when an offline converter is installed.\n")


def package_files(resume_docx: bytes, resume_pdf: bytes | None,
                  letter_docx: bytes | None, letter_pdf: bytes | None,
                  audit: dict, csv_text: str, resume_md: str = "",
                  letter_md: str = "") -> dict[str, bytes]:
    import json
    files: dict[str, bytes] = {"Tailored_Resume.docx": resume_docx}
    if resume_pdf:
        files["Tailored_Resume.pdf"] = resume_pdf
    if resume_md:
        files["Tailored_Resume.txt"] = to_text(resume_md).encode("utf-8")
        files["Tailored_Resume.md"] = resume_md.encode("utf-8")
    if letter_docx:
        files["Tailored_Cover_Letter.docx"] = letter_docx
    if letter_pdf:
        files["Tailored_Cover_Letter.pdf"] = letter_pdf
    if letter_md:
        files["Tailored_Cover_Letter.txt"] = to_text(letter_md).encode("utf-8")
    files["JD_Requirement_Audit.csv"] = csv_text.encode("utf-8-sig")
    files["Application_Match_Report.json"] = json.dumps(audit, ensure_ascii=False, indent=2).encode("utf-8")
    files["READ_ME_FIRST.txt"] = README_TEXT.encode("utf-8")
    return files


def application_zip(resume_docx: bytes, resume_pdf: bytes | None,
                    letter_docx: bytes | None, letter_pdf: bytes | None,
                    audit: dict, csv_text: str, resume_md: str = "",
                    letter_md: str = "") -> bytes:
    """Package all generated files without external storage or network."""
    return zip_files(package_files(resume_docx, resume_pdf, letter_docx, letter_pdf,
                                   audit, csv_text, resume_md, letter_md))


def zip_files(files: dict[str, bytes]) -> bytes:
    from zipfile import ZIP_DEFLATED, ZipFile

    stream = io.BytesIO()
    with ZipFile(stream, mode="w", compression=ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return stream.getvalue()


def save_package(base_dir: str | Path, label: str, files: dict[str, bytes]) -> Path:
    """Write the package to a new, timestamped local folder and return its path."""
    from datetime import datetime

    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", label).strip("_")[:60] or "application"
    folder = Path(base_dir) / f"{datetime.now():%Y-%m-%d_%H%M%S}_{safe}"
    folder.mkdir(parents=True, exist_ok=False)
    for name, data in files.items():
        (folder / name).write_bytes(data)
    return folder
