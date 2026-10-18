"""Deterministic text helpers: ingestion, sections, chunking, matching."""
import io
import unittest

from docx import Document

from fakes import CV
from textutils import (
    chunk_text, contains_phrase, format_verbatim, heading_kind, novel_numbers, proper_nouns,
    read_file, split_sections, technical_terms,
)


class MatchingTests(unittest.TestCase):
    def test_phrases(self):
        self.assertTrue(contains_phrase("Wrote C++ and Python", "C++"))
        self.assertTrue(contains_phrase("finite   element\nanalysis", "finite element analysis"))
        self.assertFalse(contains_phrase("Pythonic", "Python"))

    def test_numbers(self):
        self.assertEqual(novel_numbers("Budget 1,200 and 40%", "Budget 1200, 40 and 99%"), ["99%"])

    def test_technical_terms_and_names(self):
        self.assertEqual(technical_terms("Used MATLAB, C++, C#, Python3 and plain words in 2021"),
                         ["MATLAB", "C++", "C#", "Python3"])
        self.assertEqual(proper_nouns("Led tests. Used Kubernetes with the Senior team.\n- Built rigs"),
                         ["Kubernetes", "Senior"])
        self.assertEqual(proper_nouns("- **Programming:** Python, Abaqus"), ["Abaqus"])


class SectionTests(unittest.TestCase):
    def test_headings(self):
        self.assertEqual(heading_kind("EXPERIENCE"), "experience")
        self.assertEqual(heading_kind("## Selected Publications"), "verbatim")
        self.assertEqual(heading_kind("Technical Skills:"), "skills")
        self.assertIsNone(heading_kind("SKILLS: Python, MATLAB"))
        self.assertIsNone(heading_kind("Performed finite element analysis"))

    def test_split_sample(self):
        sections = split_sections(CV)
        self.assertEqual([s.kind for s in sections],
                         ["header", "experience", "education", "skills", "verbatim"])
        self.assertTrue(sections[0].body.startswith("A. Example"))
        self.assertFalse(sections[2].rewritable)  # education is preserved exactly
        self.assertFalse(sections[4].rewritable)

    def test_unstructured_fallback(self):
        self.assertEqual(split_sections("Jane\nDid things\nMore things")[0].kind, "document")

    def test_verbatim_format(self):
        self.assertEqual(format_verbatim("• One\n  ● Two\no Three\n# Hash"), "- One\n  - Two\n  - Three\n\\# Hash")


class ChunkTests(unittest.TestCase):
    def test_chunks_respect_size_and_keep_all_text(self):
        text = "\n\n".join(f"Paragraph {i} " + "word " * 40 for i in range(60))
        chunks = chunk_text(text, 1500)
        self.assertGreater(len(chunks), 5)
        self.assertTrue(all(len(c) <= 1500 for c in chunks))
        for i in range(60):
            self.assertTrue(any(f"Paragraph {i} " in c for c in chunks))

    def test_huge_single_line(self):
        chunks = chunk_text("x" * 5000, 1000)
        self.assertEqual("".join(chunks), "x" * 5000)


class ReadFileTests(unittest.TestCase):
    def test_txt_and_docx(self):
        self.assertIn("Alpha Motors", read_file("cv.txt", CV.encode("utf-8-sig")))
        doc = Document()
        doc.add_paragraph("Jane Doe")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Python"
        table.cell(0, 1).text = "Finite element analysis for durability studies"
        doc.add_paragraph("Final paragraph with enough text to pass the minimum length check.")
        buf = io.BytesIO()
        doc.save(buf)
        text = read_file("cv.docx", buf.getvalue())
        self.assertLess(text.index("Jane Doe"), text.index("Python | Finite"))
        self.assertLess(text.index("Python | Finite"), text.index("Final paragraph"))

    def test_rejects_unknown_and_empty(self):
        with self.assertRaises(ValueError):
            read_file("cv.rtf", b"abc")
        with self.assertRaises(ValueError):
            read_file("cv.txt", b"short")



def two_column_pdf(sidebar: list[str], main: list[str], footer: str = "Page 1 of 1") -> bytes:
    """A minimal LinkedIn-style PDF: narrow sidebar at x=22, main column at x=224."""
    def esc(t):
        return t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    ops = []
    for x, lines in ((22, sidebar), (224, main)):
        y = 740
        for line in lines:
            if line == "":
                y -= 10  # paragraph gap
                continue
            ops.append(f"BT /F1 10 Tf {x} {y} Td ({esc(line)}) Tj ET")
            y -= 13
    ops.append(f"BT /F1 9 Tf 260 30 Td ({footer}) Tj ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


SIDEBAR = ["Contact", "jane@example.test", "www.linkedin.com/in/jane-", "doe-phd (LinkedIn)", "",
           "Top Skills", "Finite Element Analysis", "", "Robot Kinematics", "", "Languages", "English"]
MAIN = ["Jane Doe,", "", "Ph.D.", "Principal Engineer at Example Robotics", "", "Experience", "",
        "Example Robotics", "Principal Engineer", "March 2019 - Present",
        "- Designed precision positioning stages and verified them with finite element analysis for",
        "thermal drift and vibration, working closely with controls engineers on interfaces.",
        "", "Education", "Doctor of Philosophy - PhD, Mechanical Engineering"]


class LayoutTests(unittest.TestCase):
    def test_two_column_pdf_reads_main_column_first(self):
        text = read_file("profile.pdf", two_column_pdf(SIDEBAR, MAIN))
        self.assertTrue(text.startswith("Jane Doe,"), text[:200])
        self.assertLess(text.index("Experience"), text.index("Contact"))
        self.assertNotIn("Page 1 of 1", text)
        # a bullet wrapped by the column width is re-joined into one line
        self.assertIn("finite element analysis for thermal drift and vibration", text)
        self.assertIn("Finite Element Analysis\n\nRobot Kinematics", text)  # list items stay separate

    def test_header_skips_contact_label_and_keeps_degree(self):
        from engine import _plan, _assemble
        from textutils import Budget
        text = read_file("profile.pdf", two_column_pdf(SIDEBAR, MAIN))
        name, contacts, sections, units = _plan(text, Budget())
        self.assertEqual(name, "Jane Doe, Ph.D.")
        self.assertIn("www.linkedin.com/in/jane-doe-phd (LinkedIn)", contacts)
        md = _assemble(name, contacts, sections, units)
        self.assertTrue(md.startswith("# Jane Doe, Ph.D.\n"))
        self.assertNotIn("## Contact", md)

    def test_split_header_plain(self):
        from textutils import split_header
        self.assertEqual(split_header("CONTACT\nJohn Roe\nPhD\njohn@x.test")[:2],
                         ("John Roe, PhD", ["john@x.test"]))


class QuoteTests(unittest.TestCase):
    SOURCE = ("Led programs at Lam Research, combining characterization, FEA, and cross-\n"
              "institutional collaboration.\nDoctor of Philosophy - PhD, Mechanical Engineering/\n"
              "Mechanical Technology\n- Developed physics-based FEA models (Thermal-\n"
              "Electrochemical) to study film-substrate interactions, reducing trials.")

    def test_layout_insensitive_quotes(self):
        from textutils import exact_quote_in
        for quote in ["combining characterization, FEA, and cross-institutional collaboration",
                      "Doctor of Philosophy – PhD, Mechanical Engineering/Mechanical Technology",
                      "“Developed physics-based FEA models (Thermal-Electrochemical)”",
                      "Led programs at Lam Research … reducing trials"]:
            self.assertTrue(exact_quote_in(quote, self.SOURCE), quote)
        self.assertFalse(exact_quote_in("Led robotics programs at Lam Research", self.SOURCE))

    def test_locate_snaps_near_quotes_and_rejects_inventions(self):
        from textutils import locate_quote
        found = locate_quote("Developed physics based FEA model (Thermal Electrochemical) to study "
                             "film substrate interactions", self.SOURCE)
        self.assertTrue(found.startswith("- Developed physics-based FEA models"), found)
        self.assertNotIn("Mechanical Technology", found)  # tightest window, not a neighbour line
        self.assertEqual(locate_quote("Designed robotic arms for wafer handling in cleanrooms",
                                      self.SOURCE), "")


CLASSIC_CV = """JANE DOE, Ph.D.
Email: jane@example.test | Mobile: +1-555-0100

Professional Summary

• Engineer with 10+ years in precision mechanisms.

Technical Skills

• Thin Film Deposition
(PVD, CVD)
• Finite Element Analysis

Professional Certification | Courses | Training

• Certification in Semiconductor Manufacturing (2025).

Work Experience

• Research Fellow | Example University, Singapore
02/2024 - 07/2025
o Project-1: Positioning stages.
▪ Designed flexure stages and verified them with FEA.

Awards and Recognition

• Best Paper award (2020).

International Publications

• Doe, J. et al. (2019). Flexure Design. Precision Eng."""


class ClassicFormatTests(unittest.TestCase):
    def test_title_case_headings_and_sub_bullets(self):
        from textutils import _clean_pdf_lines
        text = _clean_pdf_lines(CLASSIC_CV.splitlines())
        kinds = [(s.kind, s.title) for s in split_sections(text)]
        self.assertEqual([k for k, _ in kinds], ["header", "summary", "skills", "verbatim",
                                                 "experience", "verbatim", "verbatim"])
        self.assertIn("• Thin Film Deposition (PVD, CVD)", text)  # wrapped bullet re-joined
        self.assertIn("\no Project-1: Positioning stages.\n", text)  # 'o' sub-bullet kept separate

    def test_job_title_word_is_not_a_heading(self):
        self.assertIsNone(heading_kind("Research Fellow", isolated=True))
        self.assertIsNone(heading_kind("Training Manager", isolated=True))
        self.assertEqual(heading_kind("International Publications", isolated=True), "verbatim")
        self.assertIsNone(heading_kind("International Publications", isolated=False))

    def test_role_lines_and_dates(self):
        from textutils import is_date_line, role_lines
        self.assertTrue(is_date_line("02/2024 - 07/2025"))
        self.assertTrue(is_date_line("November 2025 - Present (1 year)"))
        self.assertFalse(is_date_line("Best Paper award (2020)."))
        body = "- Research Fellow | Example University, Singapore\n02/2024 - 07/2025\n- Built rigs in 2024."
        self.assertEqual(role_lines(body), ["Research Fellow | Example University, Singapore"])
        linkedin = "Example Robotics\nPrincipal Engineer\nMarch 2019 - Present (5 years)\nBoston"
        self.assertEqual(role_lines(linkedin), ["Example Robotics", "Principal Engineer"])


class WordFormatTests(unittest.TestCase):
    def test_numbered_bullets_bold_headings_nesting_and_text_boxes(self):
        from docx.enum.text import WD_BREAK  # noqa: F401  (python-docx import check)
        doc = Document()
        doc.add_paragraph("JANE DOE")
        heading = doc.add_paragraph()
        heading.add_run("Awards and Recognition").bold = True
        item = doc.add_paragraph("Best Paper award (2020).", style="List Bullet")
        sub = doc.add_paragraph("Second level detail.", style="List Bullet 2")
        role = doc.add_paragraph()
        role.add_run("Research Fellow").bold = True  # bold job title must stay content
        doc.add_paragraph("Final paragraph with enough words to pass the minimum length check.")
        buf = io.BytesIO()
        doc.save(buf)
        text = read_file("cv.docx", buf.getvalue())
        self.assertIn("## Awards and Recognition", text)
        self.assertIn("- Best Paper award (2020).", text)
        self.assertIn("Research Fellow", text)
        self.assertNotIn("## Research Fellow", text)
        self.assertEqual(item.style.name, "List Bullet")
        self.assertIn("Second level detail.", text)


if __name__ == "__main__":
    unittest.main()
