"""Word/PDF/ZIP export tests."""
import io
import json
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from engine import read_file
from exporters import application_zip, cover_letter_docx, package_files, save_package, to_docx, to_pdf, to_text
from fakes import LETTER

RESUME_MD = """# A. Example

example@research.test | +44 1234 567890

## Professional Summary

Mechanical R&D engineer experienced in finite element analysis (FEA).

## Skills

- **Simulation:** Abaqus, finite element analysis (FEA)

## Experience

**Mechanical R&D Engineer** — Alpha Motors | 2020-2025
- Performed finite element analysis (FEA) to study product durability.

## Publications

\\# Example, A. (2023). Fatigue of welded joints. Journal of Example Engineering.
"""


class ExportTests(unittest.TestCase):
    def test_resume_docx_round_trip(self):
        data = to_docx(RESUME_MD)
        doc = Document(io.BytesIO(data))
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts[0], "A. Example")
        self.assertIn("Experience", texts)
        self.assertIn("Mechanical R&D Engineer — Alpha Motors | 2020-2025", texts)
        self.assertIn("# Example, A. (2023). Fatigue of welded joints. Journal of Example Engineering.", texts)
        self.assertFalse(doc.tables)  # no tables: ATS friendly
        self.assertIn("Performed finite element analysis", read_file("r.docx", data))

    def test_letter_docx(self):
        doc = Document(io.BytesIO(cover_letter_docx(LETTER)))
        self.assertEqual(doc.paragraphs[0].text, "A. Example")
        self.assertLess(len(doc.paragraphs), 18)

    def test_plain_text(self):
        text = to_text(RESUME_MD)
        self.assertIn("EXPERIENCE", text)
        self.assertIn("• Performed finite element analysis", text)
        self.assertNotIn("**", text)

    def test_zip_contents(self):
        binary = application_zip(to_docx(RESUME_MD), None, cover_letter_docx(LETTER), None,
                                 {"score": 95}, "a,b\n1,2", resume_md=RESUME_MD, letter_md=LETTER)
        with zipfile.ZipFile(io.BytesIO(binary)) as z:
            names = set(z.namelist())
            self.assertTrue({"Tailored_Resume.docx", "Tailored_Resume.txt", "Tailored_Resume.md",
                             "Tailored_Cover_Letter.docx", "Tailored_Cover_Letter.txt",
                             "JD_Requirement_Audit.csv", "Application_Match_Report.json",
                             "READ_ME_FIRST.txt"} <= names)
            self.assertNotIn("Tailored_Resume.pdf", names)
            self.assertEqual(json.loads(z.read("Application_Match_Report.json"))["score"], 95)

    def test_save_package(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = save_package(tmp, "Beta Engineering / R&D", {"a.txt": b"hi"})
            self.assertTrue((folder / "a.txt").exists())
            self.assertIn("Beta_Engineering_R_D", folder.name)

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice not installed")
    def test_pdf_has_selectable_text_in_order(self):
        pdf = to_pdf(to_docx(RESUME_MD))
        self.assertEqual(pdf[:4], b"%PDF")
        text = read_file("r.pdf", pdf)
        self.assertLess(text.index("Professional Summary"), text.index("Experience"))
        self.assertIn("Alpha Motors", text)


if __name__ == "__main__":
    unittest.main()
